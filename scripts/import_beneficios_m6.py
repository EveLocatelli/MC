#!/usr/bin/env python3
"""Importa os benefícios dos produtos de internet da planilha do M6
(Mude M6/catalogo-produtos-TEMPLATE.xlsx) para a aba "Beneficios (M6)" da
planilha do App (catalogo-ofertas-TEMPLATE.xlsx). Só LÊ o arquivo do M6.
Reexecutar sempre que a planilha do M6 mudar (a aba é recriada)."""
import copy
import openpyxl
from openpyxl.styles import Alignment

M6 = "/Users/hitss/Desktop/Mude M6/catalogo-produtos-TEMPLATE.xlsx"
APP = "/Users/hitss/Desktop/App/catalogo-ofertas-TEMPLATE.xlsx"

# ID do produto no M6 -> ID do produto no App (None = sem correspondente no App)
MAPA = {
    "internet-350": "fibra-350mega",
    "internet-600": None,          # o App não tem 600 Mega (tem 500 Mega)
    "internet-1000": "fibra-1gb",
    "internet-5000": None,         # o App não tem 5 Giga
}
URL_PREFIXO = "https://mondrian.claro.com.br/brands/app/32px-alternative/"

def txt(v):
    return "" if v is None else " ".join(str(v).replace("\xa0", " ").split())

def marca_de_url(v):
    v = txt(v)
    return v.rsplit("/", 1)[-1].replace(".svg", "") if v.startswith("http") else v

m6 = openpyxl.load_workbook(M6, data_only=True)
linhas = []  # (id_app, id_m6, origem, ordem, titulo, icone_marca, descricao)

def varre(aba, extrai):
    for r in m6[aba].iter_rows(min_row=2, values_only=True):
        pid = txt(r[0])
        if pid in MAPA:
            linhas.append((MAPA[pid] or "", pid) + extrai(r))

def ordem(v):
    try: return int(float(v))
    except (TypeError, ValueError): return 0

varre("Card_Destaques", lambda r: ("Destaque do card", ordem(r[1]),
      txt(r[3]) or txt(r[2]), txt(r[4]) or txt(r[6]), txt(r[2])))
varre("Detalhes_Beneficios", lambda r: ("Benefício (detalhes)", ordem(r[1]),
      txt(r[4]), marca_de_url(r[2]), ""))
varre("Streamings_Incluidos", lambda r: ("Streaming incluso", ordem(r[1]),
      txt(r[2]), txt(r[2]), txt(r[3])))
varre("Detalhes_Destaques", lambda r: ("Característica (detalhes)", ordem(r[1]),
      txt(r[4]), txt(r[2]), txt(r[5])))
varre("Servicos_Digitais", lambda r: ("Serviço digital", ordem(r[1]),
      txt(r[2]), marca_de_url(r[3]), txt(r[5])))

# 500 Mega não existe no M6: os benefícios são os mesmos do 350 Mega, mudando só as velocidades
# (download 500 Mbps; upload até 50 Mbps).
TROCAS_500 = [("350 Mbps", "500 Mbps"), ("Até 35 Mbps", "Até 50 Mbps")]
def troca_500(t):
    for a, b in TROCAS_500:
        t = t.replace(a, b)
    return t
for l in [l for l in linhas if l[0] == "fibra-350mega"]:
    id_app, id_m6, origem, ordem_, titulo, marca, desc = l
    linhas.append(("fibra-500mega", "internet-350 (copiado)", origem, ordem_,
                   troca_500(titulo), marca, troca_500(desc)))

app = openpyxl.load_workbook(APP)
nome = "Beneficios (M6)"
if nome in app.sheetnames:
    del app[nome]
ws = app.create_sheet(nome)
src = app["Produtos"]["A1"]
heads = [("ID do produto no App (vazio = sem correspondente)", 30), ("ID do produto no M6", 22),
         ("Origem na planilha do M6", 24), ("Ordem", 8), ("Título / texto", 44),
         ("Ícone ou marca", 28), ("Descrição / observação", 90)]
for i, (h, w) in enumerate(heads, 1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = copy.copy(src.font); c.fill = copy.copy(src.fill)
    c.alignment = copy.copy(src.alignment); c.border = copy.copy(src.border)
    ws.column_dimensions[c.column_letter].width = w
for r, row in enumerate(linhas, 2):
    for i, v in enumerate(row, 1):
        ws.cell(row=r, column=i, value=v).alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
ws.row_dimensions[1].height = app["Produtos"].row_dimensions[1].height or 45
app.save(APP)
print(f"OK — {len(linhas)} linhas importadas para a aba '{nome}'")
